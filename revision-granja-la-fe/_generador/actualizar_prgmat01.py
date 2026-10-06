# -*- coding: utf-8 -*-
"""
Pasa el PRGMAT01 oficial de la v05 a la v06 con el plan nuevo de G-GRANJA LA FE y actualiza los
datos del Plan AM que usa IncubApp (src/data/annualMaintenancePlanData.json).
Uso (desde la raíz del repo):  python revision-granja-la-fe/_generador/actualizar_prgmat01.py
"""
import copy, json, os, sys
import openpyxl
from openpyxl.styles import PatternFill, Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from contenido_plan import TAREAS, MODULOS, TANQUES
from generar import cronograma, semanas, HOY

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
XLSX = os.path.join(ROOT, "src", "assets", "manuales", "sig", "PROCEDIMIENTOS", "PRGMAT01 Programa de mantenimiento preventivo.xlsx")
APP_JSON = os.path.join(ROOT, "src", "data", "annualMaintenancePlanData.json")
SEDE_XLSX = "G-GRANJA LA FE"
SEDE_APP = "GRANJA LA FE"
VIGENTE_DESDE = 42  # semana ISO de 2026 en que entra en vigor la v06 para La Fe (12-oct-2026)

BAND, CRIT, BIO, PEND = "EDF2F8", "FBE3DD", "E4F0E4", "FFF9E6"


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def color_tarea(t, i):
    if t["tipo"].startswith("Crítico"): return CRIT
    if any(e.startswith("(") for e in t["equipos"]) and not any(not e.startswith("(") for e in t["equipos"]): return PEND
    if t["tipo"] in ("Bioseguridad", "Legal"): return BIO
    return BAND if i % 2 == 0 else None


def estilo(dst, src, fill_hex=None):
    dst.font = copy.copy(src.font); dst.border = copy.copy(src.border); dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format
    dst.fill = fill(fill_hex) if fill_hex else PatternFill(fill_type=None)


def mover_filas(ws, fila, borrar, insertar):
    """Borra «borrar» filas desde «fila» e inserta «insertar»; corre también las celdas combinadas de abajo
    (openpyxl no las mueve y al guardar borraría los valores de la fila que quede debajo)."""
    abajo = [r for r in list(ws.merged_cells.ranges) if r.min_row >= fila]
    for r in abajo:
        assert r.min_row >= fila + borrar, f"celda combinada dentro del bloque: {r}"
        ws.unmerge_cells(str(r))
    if borrar: ws.delete_rows(fila, borrar)
    if insertar: ws.insert_rows(fila, insertar)
    d = insertar - borrar
    for r in abajo:
        ws.merge_cells(start_row=r.min_row + d, start_column=r.min_col, end_row=r.max_row + d, end_column=r.max_col)


def n_equipos(t):
    n = len([e for e in t["equipos"] if not e.startswith("(")])
    return n or "—"


def main():
    wb = openpyxl.load_workbook(XLSX)

    # 00 PORTADA
    ws = wb["00 PORTADA"]
    ws["B9"] = "06  ·  OFICIAL"
    ws["B10"] = HOY

    # 01 CONTENIDO
    ws = wb["01 CONTENIDO"]
    ws["C11"] = "Equipos reales por sede: 101 de Mántum en planta + 8 sin registrar; G-GRANJA LA FE: 24 galpones y 11 tanques (silos, motobombas, tableros generales y planta se codifican en el levantamiento LF-100). Campos a diligenciar en amarillo."
    ws["C13"] = "58 tareas (LF-100 a LF-221): alimentación (silos, tolvas, canal y cadena), recolección de huevo, nidos, agua y antipercheo, iluminación inteligente, tanques, motobombas, tableros eléctricos, chumaceras, planta eléctrica. Manuales MAN-LF-01 a MAN-LF-06."

    # 02 OBJETO Y ALCANCE
    ws = wb["02 OBJETO Y ALCANCE"]
    ws["B6"] = ws["B6"].value.replace("y G-GRANJA LA FE (24 galpones),", "y G-GRANJA LA FE (24 galpones: 12 de producción con equipos Big Herdsman y 12 de levante, más sus silos, tanques, motobombas, tableros eléctricos y planta eléctrica),")

    # 05 INVENTARIO — La Fe
    ws = wb["05 INVENTARIO"]
    filas_lf = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == SEDE_XLSX]
    modelo = filas_lf[0]
    prod = {g for _, uso, gal, _ in MODULOS if uso == "Producción" for g in gal}
    for r in filas_lf:
        nombre = str(ws.cell(r, 3).value)
        num = nombre.split(" ")[1].split("-")[0] if nombre.startswith("Galpón") else ""
        es_prod = num[:1] in "234"
        ws.cell(r, 4).value = "Galpón de producción" if es_prod else "Galpón de levante"
        ws.cell(r, 5).value = "BIG HERDSMAN" if es_prod else "(confirmar en LF-100)"
    nuevos = [[SEDE_XLSX, t, f"Tanque de agua {t}", "Agua", "(LF-100)", "(LF-100)", "(LF-100)", "(LF-100)", "", "Activo"] for t in TANQUES]
    nuevos += [[SEDE_XLSX, c, n, s, "(LF-100)", "(LF-100)", "(LF-100)", "(LF-100)", "", "Por codificar en LF-100"] for c, n, s in [
        ("SI-LF-nn", "Silos de alimento", "Alimentación"), ("MB-LF-nn", "Motobombas de suministro de agua", "Agua"),
        ("TG-LF-nn", "Tablero general, transferencia y tableros de bombeo", "Energía"), ("PE-LF-01", "Planta eléctrica de emergencia + ATS", "Energía")]]
    ult = filas_lf[-1]
    mover_filas(ws, ult + 1, 0, len(nuevos))
    for i, vals in enumerate(nuevos):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(ult + 1 + i, c, v); estilo(cell, ws.cell(modelo, c), PEND if "LF-100" in str(v) else None)
    for r in range(ult + len(nuevos) + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith("TOTAL EQUIPOS"):
            ws.cell(r, 1).value = v.replace("G-GRANJA LA FE: 24 galpones", "G-GRANJA LA FE: 24 galpones + 11 tanques (silos, motobombas, tableros generales y planta: se codifican en el levantamiento LF-100)")

    # 06.2 PLAN LA FE
    ws = wb["06.2 PLAN LA FE"]
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Código")
    ref_h, ref_d = [ws.cell(hdr, c) for c in range(1, 15)], [ws.cell(hdr + 1, c) for c in range(1, 15)]
    ref_h = [copy.copy(c._style) for c in ref_h]; ref_d = [copy.copy(c._style) for c in ref_d]
    if ws.max_row > hdr: mover_filas(ws, hdr + 1, ws.max_row - hdr, 0)
    for c, titulo in ((15, "Manual"), (16, "Reemplaza a (v05)")):
        cell = ws.cell(hdr, c, titulo); cell._style = copy.copy(ref_h[0])
    ws.column_dimensions["O"].width = 20; ws.column_dimensions["P"].width = 16
    for i, t in enumerate(TAREAS):
        r = hdr + 1 + i
        vals = [t["code"], t["system"], t["eq_class"], ", ".join(t["equipos"]), n_equipos(t), t["desc"], t["tipo"], t["freq"], t["parada"], t["dur"], t["resp"], t["crit"], t["sust"], t["evid"], t["manual"], t["reemplaza"]]
        f = color_tarea(t, i)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell._style = copy.copy(ref_d[min(c, 14) - 1])
            cell.fill = fill(f) if f else PatternFill(fill_type=None)
    ws.auto_filter.ref = None

    # 07 CRONOGRAMA 52 SEM
    ws = wb["07 CRONOGRAMA 52 SEM"]
    filas = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("LF-")]
    a, b = filas[0], filas[-1]
    est_txt = [copy.copy(ws.cell(a, c)._style) for c in range(1, 7)]
    est_on = copy.copy(next(ws.cell(r, c)._style for r in filas for c in range(7, 59) if ws.cell(r, c).value == "●"))
    est_off = copy.copy(ws.cell(a, 7)._style)
    mover_filas(ws, a, b - a + 1, len(TAREAS))
    for i, t in enumerate(TAREAS):
        r = a + i
        cr = cronograma(t)
        rot = t["rot"] if cr or t["freq"] in ("Diaria",) else t["rot"] + " · se programa con la fecha real de vaciado"
        for c, v in enumerate([t["code"], SEDE_XLSX, t["eq_class"], t["desc"], t["freq"], rot], 1):
            cell = ws.cell(r, c, v); cell._style = copy.copy(est_txt[c - 1])
        for w in range(1, 53):
            cell = ws.cell(r, 6 + w, "●" if w in cr else None)
            cell._style = copy.copy(est_on if w in cr else est_off)

    # 08 CARGA — se reescriben las filas de La Fe y su total
    ws = wb["08 CARGA"]
    VECES = {"Diaria": 365, "Semanal": 52, "Mensual": 12, "Trimestral": 4, "Semestral": 2, "Anual": 1}
    agg = {}
    for t in TAREAS:
        a_ = agg.setdefault(t["freq"], [0, 0, 0.0]); n = len(t["equipos"]); v = VECES.get(t["freq"], 1)
        a_[0] += 1; a_[1] += n * v; a_[2] += n * v * t["dur"]
    lf = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == SEDE_XLSX and ws.cell(r, 2).value != "TOTAL"]
    tot_r = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == SEDE_XLSX and ws.cell(r, 2).value == "TOTAL")
    est = [copy.copy(ws.cell(lf[0], c)._style) for c in range(1, 7)]
    mover_filas(ws, lf[0], len(lf), len(agg))
    total = 0.0
    for i, (k, (nt, ni, hh)) in enumerate(sorted(agg.items())):
        nota = "Rondas del operario galponero, dentro de su jornada" if k == "Diaria" else ("≈ 1 vacío sanitario por galpón al año" if k.startswith("Por") else None)
        for c, v in enumerate([SEDE_XLSX, k, nt, ni, round(hh, 1), nota], 1):
            cell = ws.cell(lf[0] + i, c, v); cell._style = copy.copy(est[c - 1])
        total += hh
    tot_r = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == SEDE_XLSX and ws.cell(r, 2).value == "TOTAL")
    sin_rondas = total - agg.get("Diaria", [0, 0, 0])[2]
    ws.cell(tot_r, 5).value = round(total, 1)
    ws.cell(tot_r, 6).value = f"≈ {total / 1800:.1f} personas-año (base 1.800 h); sin las rondas diarias del galponero: {sin_rondas:.0f} h ≈ {sin_rondas / 1800:.1f} personas-año de mantenimiento"

    # 12 PENDIENTES
    ws = wb["12 PENDIENTES"]
    for r in range(1, ws.max_row + 1):
        v = str(ws.cell(r, 2).value or "")
        if v.startswith("Obtener los manuales OEM de Petersime y Big Herdsman"):
            ws.cell(r, 2).value = "Obtener los manuales OEM de Petersime y contrastar las periodicidades. (Big Herdsman: RESUELTO en la v06 — manuales MAN-LF-01 a MAN-LF-06 aprobados como documentos de la empresa y carta de solicitud del original al fabricante.)"
        if v.startswith("Inventariar la planta eléctrica de emergencia, el ATS, los silos y las bombas de LA FE"):
            ws.cell(r, 2).value = "RESUELTO en la v06 como tarea programada: levantamiento de campo LF-100 (semanas 42-43 de 2026) con formato prellenado por galpón — planta, ATS, silos, motobombas, tableros, chumaceras, placas de motores y equipo de levante."

    # 13 CONTROL DE CAMBIOS
    ws = wb["13 CONTROL DE CAMBIOS"]
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value)
    r = last + 1
    texto = ("PLAN DE G-GRANJA LA FE REHECHO SOBRE LOS EQUIPOS REALES (aprobado por el Área de Mantenimiento el 2026-10-06). "
             "(1) La hoja 06.2 pasa de 30 a 58 tareas, con códigos LF-100 a LF-221; los LF-001 a LF-030 quedan retirados (columna «Reemplaza a» para trazabilidad). "
             "(2) La alimentación de producción se corrige a canal y cadena con tolvas y silos (Big Herdsman); se separan las tareas de levante. "
             "(3) Se agregan los sistemas que faltaban: antipercheo, iluminación inteligente, tableros eléctricos, motobombas, chumaceras e inspección estructural de tanques; limpieza de silos y overhaul por vacío sanitario. "
             "(4) Cronograma por módulo en ciclo de 4 semanas (M2, M3, M4, levante) por bioseguridad y carga; vigente para La Fe desde la semana 42 de 2026. "
             "(5) Manuales MAN-LF-01 a MAN-LF-06 e indicaciones de cada actividad (pasos PROMAT01 + pasos técnicos). "
             "(6) Los datos de placa y equipos sin código se levantan con la tarea única LF-100 (semanas 42-43). (7) Estándares adoptados: antipercheo ≥ 3 kV al final de línea; grasa NSF H1; aceite de reductor cada 5.000 h o 2 años sin otro dato.")
    for c, v in enumerate(["06", HOY, texto, "Área de Mantenimiento"], 1):
        cell = ws.cell(r, c, v); cell._style = copy.copy(ws.cell(last, c)._style)

    wb.save(XLSX)
    print("PRGMAT01 actualizado a v06:", XLSX)
    actualizar_app()


def actualizar_app():
    """Reemplaza las tareas de La Fe en los datos del Plan AM de IncubApp (idempotente)."""
    data = json.load(open(APP_JSON, encoding="utf-8"))
    otras = [t for t in data["tasks"] if t.get("sede") != SEDE_APP]
    nuevas = []
    for t in TAREAS:
        eq = [e for e in t["equipos"] if not e.startswith("(")]
        nuevas.append({
            "code": t["code"], "sede": SEDE_APP, "system": t["system"].upper(), "equipmentClass": t["eq_class"],
            "applyingEquipment": ", ".join(t["equipos"]), "count": len(eq) or 1, "description": t["desc"], "type": t["tipo"],
            "frequency": t["freq"], "shutdown": t["parada"], "duration": str(t["dur"]), "responsible": t["resp"],
            "acceptanceCriteria": t["crit"], "frequencyRationale": t["sust"], "evidenceFormat": t["evid"],
            "isCriticalSecurity": t["cs"], "isBiosecurity": t["bio"],
            "cronograma": {"rotacion": t["rot"], "weeks": [w for w in semanas(t) if w >= VIGENTE_DESDE], "vigenteDesdeSemana": VIGENTE_DESDE},
            "manual": t["manual"], "replaces": t["reemplaza"],
            "technicalSteps": t["pasos"], "safety": t["seg"], "tools": t["herr"],
        })
    # Se conserva el orden: planta, La Fe, La Esperanza
    idx = next((i for i, t in enumerate(data["tasks"]) if t.get("sede") == SEDE_APP), len(otras))
    tareas = [t for t in data["tasks"][:idx] if t.get("sede") != SEDE_APP] + nuevas + [t for t in data["tasks"][idx:] if t.get("sede") != SEDE_APP]
    sistemas = {}
    for t in tareas:
        s = sistemas.setdefault(t["system"], {"name": t["system"], "sedes": [], "taskCount": 0, "equipmentClasses": [], "criticalCount": 0, "biosecurityCount": 0})
        if t["sede"] not in s["sedes"]: s["sedes"].append(t["sede"])
        s["taskCount"] += 1
        if t["equipmentClass"] and t["equipmentClass"] not in s["equipmentClasses"]: s["equipmentClasses"].append(t["equipmentClass"])
        s["criticalCount"] += bool(t.get("isCriticalSecurity")); s["biosecurityCount"] += bool(t.get("isBiosecurity"))
    systems = [{"name": s["name"], "sedes": s["sedes"], "taskCount": s["taskCount"], "equipmentCount": len(s["equipmentClasses"]),
                "equipmentClasses": s["equipmentClasses"], "criticalCount": s["criticalCount"], "biosecurityCount": s["biosecurityCount"]} for s in sistemas.values()]
    data["tasks"] = tareas
    data["systems"] = systems
    data["metadata"].update({"generatedAt": HOY + "T12:00:00.000Z", "source": "PRGMAT01 v06 & SIG-MANTENIMIENTO REGISTROS",
                             "totalTasks": len(tareas), "totalSystems": len(systems)})
    with open(APP_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2); f.write("\n")
    print(f"IncubApp: {len(tareas)} tareas ({len(nuevas)} de La Fe), {len(systems)} sistemas")


if __name__ == "__main__":
    # --solo-app: el PRGMAT01 ya está en v06; solo se recargan los datos de la app.
    actualizar_app() if "--solo-app" in sys.argv else main()
