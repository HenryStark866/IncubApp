# -*- coding: utf-8 -*-
"""
Genera los documentos del Plan AM de G-GRANJA LA FE (v06 aprobado):
  Manuales/MAN-LF-0X *.docx · PRGMAT01 06.2 Plan AM Granja La Fe (v06 aprobado).xlsx
  Indicaciones de cada actividad - Plan AM Granja La Fe.docx · datos/plan_am_la_fe.json
Uso:  python generar.py   (requiere python-docx y openpyxl)
"""
import json, os, re, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from contenido_plan import TAREAS, MODULOS, MANUALES as MAN_TIT, GRUPOS, pasos_promat, PROD_M2, PROD_M3, PROD_M4, LEV_M1, LEV_M5, TANQUES
from contenido_manuales import MANUALES, AVISO_COMUN, REGISTROS

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HOY = datetime.date(2026, 10, 6).isoformat()
EMPRESA = "ANTIOQUEÑA DE INCUBACIÓN S.A.S."
NIT = "900.762.687"
VERSION_DOC = "01 · APROBADO 2026-10-06"
NAVY = "0B1428"; ORANGE = "E0740A"; GRIS = "F3F6FA"

# ───────────────────────── Cronograma 52 semanas ─────────────────────────
# Regla de bioseguridad y de carga: un módulo por semana en ciclo de 4 semanas
#   S1 → M2 · S2 → M3 · S3 → M4 · S4 → levante (M1 y M5)
# Así todas las tareas mensuales/trimestrales/semestrales de un módulo se hacen la misma semana
# y el técnico no cruza entre módulos de producción en la misma semana.
CICLO = ["M2", "M3", "M4", "LEV"]
def modulo_de(g):
    if g in PROD_M2: return "M2"
    if g in PROD_M3: return "M3"
    if g in PROD_M4: return "M4"
    if g in LEV_M1 or g in LEV_M5: return "LEV"
    return None

PERIODO = {"Mensual": 4, "Trimestral": 13, "Semestral": 26, "Anual": 52}

def lotes_de(t):
    mods = []
    for g in t["equipos"]:
        m = modulo_de(g)
        if m and m not in mods: mods.append(m)
    return mods  # vacío → equipo único / sin módulo

def cronograma(t):
    """Devuelve {semana: etiqueta} para la tarea."""
    f = t["freq"]
    if f.startswith("Única vez"):
        return {42: "●", 43: "●"}  # semanas 42-43 de 2026 (12 al 25 de octubre)
    if f == "Semanal":
        return {w: "●" for w in range(1, 53)}
    if f not in PERIODO:
        return {}  # Diaria (ronda) o por vacío sanitario / por ciclo: se fija con la fecha real de vaciado
    p = PERIODO[f]
    mods = lotes_de(t)
    out = {}
    if not mods:  # equipo único o grupo sin módulo: semana 2 del periodo
        start = 2 if p > 4 else 1
        for k in range(0, 52, p):
            if start + k <= 52: out[start + k] = "●"
        return out
    for m in mods:
        pos = CICLO.index(m) + 1  # semana del ciclo de 4 en la que se visita el módulo
        k = 0
        while True:
            objetivo = 1 + k * p
            if objetivo > 52: break
            w = objetivo
            while (w - 1) % 4 + 1 != pos: w += 1
            if w <= 52: out[w] = (out.get(w, "●") + ("2" if m == "M2" else "3" if m == "M3" else "4" if m == "M4" else "L"))
            k += 1
    return out

def semanas(t):
    return sorted(cronograma(t).keys())

# ───────────────────────── Utilidades docx ─────────────────────────
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_)
    tcPr.append(s)

def set_base(doc):
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)

def membrete(doc, codigo, titulo, version=VERSION_DOC):
    t = doc.add_table(rows=3, cols=3); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a = t.cell(0, 0).merge(t.cell(2, 0)); a.text = ""
    p = a.paragraphs[0]; r = p.add_run(EMPRESA); r.bold = True; r.font.size = Pt(9)
    p.add_run(f"\nNIT {NIT}\nHispania, Antioquia").font.size = Pt(8)
    b = t.cell(0, 1).merge(t.cell(1, 1)); b.text = ""
    p = b.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(NAVY)
    t.cell(2, 1).text = "GESTIÓN DE MANTENIMIENTO · G-GRANJA LA FE"
    t.cell(2, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (k, v) in enumerate([("Código", codigo), ("Versión", version), ("Fecha", HOY)]):
        c = t.cell(i, 2); c.text = ""; rr = c.paragraphs[0].add_run(f"{k}: "); rr.bold = True; rr.font.size = Pt(8)
        c.paragraphs[0].add_run(v).font.size = Pt(8)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                for rr in pp.runs:
                    if rr.font.size is None: rr.font.size = Pt(8)
    doc.add_paragraph()

def h(doc, txt, lvl=1):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True
    r.font.size = Pt(13 if lvl == 1 else 11); r.font.color.rgb = RGBColor.from_string(NAVY if lvl == 1 else ORANGE)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p

def aviso(doc, txt):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.cell(0, 0); shade(c, "FFF8EF"); c.text = ""
    r = c.paragraphs[0].add_run(txt); r.font.size = Pt(9)
    doc.add_paragraph()

def tabla(doc, headers, rows, widths=None, font=8.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; shade(c, NAVY)
        r = c.paragraphs[0].add_run(hd); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(font)
    # encabezado repetido
    trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

def lista(doc, items, numerada=False):
    for it in items:
        doc.add_paragraph(it, style="List Number" if numerada else "List Bullet")

def firmas(doc):
    h(doc, "Aprobación", 2)
    tabla(doc, ["", "Nombre", "Cargo", "Firma", "Fecha"], [["Elaboró — Área de Mantenimiento", "", "", "", ""], ["Revisó — Coordinación de Granja", "", "", "", ""], ["Aprobó — Gerencia", "", "", "", ""]], font=9)

def resumen_equipos(eq):
    extra = [e for e in eq if e.startswith("(")]
    base = [e for e in eq if not e.startswith("(")]
    if extra and base: return resumen_equipos(base) + " + " + " ".join(extra)
    if extra: return " ".join(extra)
    if eq == TANQUES: return "11 tanques: " + ", ".join(eq)
    if len(eq) == 24: return "Los 24 galpones (levante y producción)"
    if eq == PROD_M2 + PROD_M3 + PROD_M4: return "12 galpones de producción (G201–G204, G301–G304, G401–G404)"
    if eq == LEV_M1 + LEV_M5: return "12 galpones de levante (G101–G106, G501–G506)"
    return ", ".join(eq)

# ───────────────────────── Manuales ─────────────────────────
def generar_manual(m):
    doc = Document(); set_base(doc)
    membrete(doc, m["code"], m["titulo"])
    p = doc.add_paragraph(); r = p.add_run(m["subtitulo"]); r.italic = True; r.font.size = Pt(11)
    p = doc.add_paragraph(); r = p.add_run("Fabricante: "); r.bold = True
    p.add_run("Qingdao Big Herdsman Machinery Co., Ltd. (Qingdao, Shandong, China) — bigherdsman.com")
    h(doc, "1. Objeto y alcance")
    doc.add_paragraph(f"Establecer cómo se opera, inspecciona, mantiene, limpia y diagnostica el equipo, en concordancia con el Plan AM de la granja (PRGMAT01 hoja 06.2) y el procedimiento PROMAT01. Alcance: {m['alcance']}")
    h(doc, "2. Sobre este documento")
    aviso(doc, AVISO_COMUN)
    for titulo, bloques in m["secciones"]:
        h(doc, titulo)
        for b in bloques:
            k = b[0]
            if k == "p": doc.add_paragraph(b[1])
            elif k == "ul": lista(doc, b[1])
            elif k == "ol": lista(doc, b[1], True)
            elif k == "aviso": aviso(doc, b[1])
            elif k == "tabla": tabla(doc, b[1], b[2])
            elif k == "plan":
                rows = [[t["code"], t["desc"], t["freq"], t["resp"], t["crit"]] for t in TAREAS if t["grupo"] == b[1]]
                tabla(doc, ["Código", "Actividad", "Frecuencia", "Responsable", "Criterio de aceptación"], rows, widths=[2, 6.5, 2.2, 2.3, 4], font=8)
    n = len(m["secciones"]) + 3
    h(doc, f"{n}. Registros")
    tabla(doc, ["Registro", "Diligencia", "Cuándo"], REGISTROS)
    doc.add_paragraph("Los registros se conservan mínimo un año en la granja (Res. ICA 3651 de 2014 Art. 4.2.13). El soporte primario es IncubApp / Mántum; lo diligenciado en físico se digitaliza al cerrar la orden.")
    h(doc, f"{n + 1}. Datos de placa y manual original")
    lista(doc, [
        "Los datos marcados con ⚠ se toman en el levantamiento de campo LF-100 (semanas 42-43) y se registran en la hoja 05 INVENTARIO del PRGMAT01; desde ese momento son el valor de referencia de este manual.",
        "Lubricantes y cambios de aceite de reductores: los de la placa del reductor; mientras no haya otro dato, aceite mineral cada 5.000 horas de trabajo o cada 2 años, lo que ocurra primero (estándar de la empresa).",
        "Horarios, alturas, presiones e intensidades de luz: los del programa vigente de Producción para el lote; este manual no fija esos valores.",
        "Solicitud del manual original enviada al fabricante (ver carta en la carpeta del plan). Si llega, se contrasta y se emite la versión 02.",
    ])
    firmas(doc)
    ruta = os.path.join(BASE, "Manuales", f"{m['code']} {MAN_TIT[m['code']]}.docx".replace(":", " -"))
    doc.save(ruta); return ruta

# ───────────────────────── Indicaciones de cada actividad ─────────────────────────
def generar_indicaciones():
    doc = Document(); set_base(doc)
    membrete(doc, "PRGMAT01-06.2-IND", "INDICACIONES DE CADA ACTIVIDAD — PLAN AM G-GRANJA LA FE")
    doc.add_paragraph("Cada actividad del plan se ejecuta con los mismos seis pasos del procedimiento PROMAT01 v02 sección 4 que IncubApp muestra al abrir una actividad del Plan AM de Planta Incubant (programar, ejecutar, verificar, validar bioseguridad, registrar y liberar). A diferencia de la planta, aquí se agregan los pasos técnicos de cada tarea, sacados del manual del equipo (MAN-LF-01 a MAN-LF-06).")
    aviso(doc, "Documento aprobado (versión 01, 2026-10-06), parte del PRGMAT01 v06. Los pasos técnicos salen de los manuales MAN-LF-01 a MAN-LF-06; los datos de placa se completan con el levantamiento LF-100. Los registros se conservan mínimo un año (Res. ICA 3651/2014 Art. 4.2.13).")
    sistema_actual = None
    for t in TAREAS:
        if t["system"] != sistema_actual:
            sistema_actual = t["system"]
            doc.add_page_break(); h(doc, sistema_actual.upper())
        h(doc, f"{t['code']} · {t['desc']}", 2)
        datos = [
            ("Clase de equipo", t["eq_class"]), ("Equipos que aplican", resumen_equipos(t["equipos"])),
            ("Tipo / frecuencia", f"{t['tipo']} · {t['freq']}"), ("¿Parada?", t["parada"]), ("Duración estimada", f"{t['dur']} h por equipo"),
            ("Responsable", t["resp"]), ("Criterio de aceptación", t["crit"]), ("Registro / evidencia", t["evid"]),
            ("Manual de referencia", t["manual"]), ("Sustento", t["sust"]),
        ]
        if t["cs"]: datos.append(("Marca", "TAREA CRÍTICA DE SEGURIDAD"))
        tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
        for k, v in datos:
            c = tb.add_row().cells; c[0].text = ""; shade(c[0], GRIS)
            r = c[0].paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(8.5)
            c[1].text = ""; r = c[1].paragraphs[0].add_run(v); r.font.size = Pt(8.5)
            if k == "Marca": r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
            c[0].width = Cm(4); c[1].width = Cm(13)
        doc.add_paragraph()
        p = doc.add_paragraph(); r = p.add_run("Pasos de ejecución (PROMAT01 v02 · sección 4)"); r.bold = True
        rows = []
        for titulo, det, resp, reg in pasos_promat(t):
            rows.append([titulo, " ".join(det), resp, reg])
        tabla(doc, ["Paso", "Qué hacer", "Responsable", "Registro"], rows, widths=[3, 8, 3, 3.5], font=8)
        p = doc.add_paragraph(); r = p.add_run("Pasos técnicos del paso 2 (según el manual del equipo)"); r.bold = True
        lista(doc, t["pasos"], True)
        if t["seg"]:
            p = doc.add_paragraph(); r = p.add_run("Seguridad"); r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
            lista(doc, t["seg"])
        if t["herr"]:
            p = doc.add_paragraph(); r = p.add_run("Herramientas y materiales: "); r.bold = True
            p.add_run(", ".join(t["herr"]) + ".")
    ruta = os.path.join(BASE, "Indicaciones de cada actividad - Plan AM Granja La Fe.docx")
    doc.save(ruta); return ruta


# ───────────────────────── Hoja LEVANTAMIENTO (LF-100) ─────────────────────────
EQUIPOS_PROD = [
    ("Alimentación", "Motorreductor unidad motriz canal y cadena"), ("Alimentación", "Motor transportador silo-tolva"),
    ("Alimentación", "Sensor de nivel de tolva"), ("Alimentación", "Comedero de machos (¿existe?)"),
    ("Recolección", "Motorreductor banda de huevo"), ("Nidos", "Motorreductor apertura/cierre de nidos"), ("Nidos", "Controlador de nidos"),
    ("Antipercheo", "Energizador"), ("Iluminación", "Controlador y dimmer de luz"), ("Iluminación", "Luminarias LED (cantidad, W, IP)"),
    ("Chumaceras", "Chumaceras banda (cantidad y referencia)"), ("Chumaceras", "Chumaceras nidos (cantidad y referencia)"),
    ("Chumaceras", "Chumaceras alimentación/transportador (cantidad y referencia)"), ("Energía", "Tablero del galpón"),
]
EQUIPOS_LEV = [
    ("Alimentación", "Comedero de levante: tipo, marca y motor"), ("Agua", "Bebedero: tipo y marca"), ("Iluminación", "Controlador y luminarias"),
    ("Antipercheo", "¿Tiene? Energizador"), ("Energía", "Tablero del galpón"),
]
EQUIPOS_GRANJA = [
    ("Alimentación", "Silo (SI-LF-nn): capacidad y galpón que alimenta"), ("Agua", "Motobomba (MB-LF-nn)"), ("Agua", "Tanque hidroneumático / presostato"),
    ("Energía", "Tablero general (TG-LF-01)"), ("Energía", "Tablero de bombeo (TG-LF-nn)"), ("Energía", "Planta eléctrica (PE-LF-01)"), ("Energía", "Transferencia ATS"),
] + [("Agua", f"Tanque {t}: material, capacidad, tipo de base") for t in TANQUES]


def hoja_levantamiento(wb):
    ws = wb.create_sheet("LEVANTAMIENTO LF-100")
    cols = ["Galpón / ubicación", "Sistema", "Equipo", "Código asignado", "Marca", "Modelo", "Potencia (kW/HP)", "Tensión / Hz", "In (A)", "N.º serie", "Referencia / cantidad", "Foto (sí/no)", "Observación"]
    titulo(ws, "LEVANTAMIENTO DE CAMPO — TAREA LF-100 (semanas 42-43)", "Una fila por equipo esperado. Diligenciar las celdas amarillas en campo, fotografiar cada placa y trasladar los datos a la hoja 05 INVENTARIO del PRGMAT01 y a IncubApp. Si un equipo no existe, escribir «No existe».", len(cols))
    hdr(ws, 5, cols)
    r = 6
    filas = []
    for mod, uso, gal, _ in MODULOS:
        for g in gal:
            for sis, eq in (EQUIPOS_PROD if uso == "Producción" else EQUIPOS_LEV):
                filas.append((f"{g} · {mod}", sis, eq))
    for sis, eq in EQUIPOS_GRANJA:
        filas.append(("Granja", sis, eq))
    for a, b, c in filas:
        celda(ws, r, 1, a); celda(ws, r, 2, b); celda(ws, r, 3, c)
        cod = a.split(" ")[0] if eq_es_tablero(c) else None
        for k in range(4, len(cols) + 1):
            celda(ws, r, k, cod if (k == 4 and cod) else None, fill="FFF4C2")
        r += 1
    for i, w in enumerate([18, 13, 44, 14, 14, 14, 12, 12, 8, 14, 22, 8, 28], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D6"; ws.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{r - 1}"


def eq_es_tablero(nombre):
    return nombre.startswith("Tablero del galpón")


# ───────────────────────── Hoja PENDIENTES (estado al 2026-10-06) ─────────────────────────
PENDIENTES = [
    ("Manuales originales de Big Herdsman.", "RESUELTO", "Big Herdsman no publica manuales. Se aprobaron como documentos de la empresa los manuales MAN-LF-01 a MAN-LF-06 (v01). Queda redactada la carta de solicitud del manual original al fabricante (carpeta del plan); si llega, se contrasta y se emite la v02. No bloquea: la Res. ICA 3651 exige que la empresa defina y documente el procedimiento, y ya está hecho.", "Mantenimiento — enviar la carta"),
    ("Datos de placa de los motores (alimentación, banda, nidos, transportador, motobombas).", "PROGRAMADO", "Convertido en la tarea LF-100 (semanas 42-43) con la hoja «LEVANTAMIENTO LF-100» prellenada por galpón. Mientras tanto, cada OT con medición de corriente fotografía y registra la placa del motor que mide (primer paso de LF-103, LF-123, LF-133, LF-162).", "Téc. + Aux. mantenimiento — semanas 42-43"),
    ("Equipo de levante (módulos 1 y 5).", "PROGRAMADO", "Incluido en LF-100. Las tareas LF-109 a LF-111, LF-141 a LF-144 y LF-181 a LF-185 aplican a cualquier línea automática y se ejecutan desde ya; con el dato se precisa la clase de equipo.", "Supervisión de granja — semana 42"),
    ("Silos, motobombas, tableros generales, planta y ATS sin código.", "PROGRAMADO", "Esquema de códigos definido: SI-LF-nn, MB-LF-nn, TG-LF-nn, PE-LF-01. Se asignan y etiquetan en LF-100. Las tareas ya se ejecutan identificando el equipo en la OT.", "Mantenimiento — semanas 42-43"),
    ("Comederos separados para machos.", "PROGRAMADO", "Pregunta incluida en LF-100 (fila «Comedero de machos» de cada galpón de producción). Si existen, aplican las mismas tareas LF-101 a LF-108.", "Producción — semana 42"),
    ("Valores de programa (luz, horario de nidos, presión de agua por edad, meta de huevo roto).", "RESUELTO", "Resuelto por diseño: los criterios remiten al programa vigente del lote que publica Producción. El plan no fija esos valores para no tener que cambiarse con cada lote.", "—"),
    ("Tensión mínima del energizador antipercheo.", "RESUELTO", "Estándar adoptado por la empresa: ≥ 3 kV al final de cada línea (LF-172). Si la placa o el manual del energizador exigen más, prevalece ese valor.", "—"),
    ("Aceite de reductores y lubricantes.", "RESUELTO", "Grasa grado alimenticio NSF H1 para chumaceras y mecanismos en contacto con huevo o alimento; aceite del reductor según su placa y, sin otro dato, cambio cada 5.000 h de trabajo o cada 2 años (estándar de la empresa).", "—"),
    ("Cargar el plan en IncubApp y reemplazar la hoja 06.2 del PRGMAT01.", "RESUELTO", "PRGMAT01 actualizado a v06 (hojas 00, 01, 05, 06.2, 07, 08, 12 y 13) y plan cargado en IncubApp (src/data/annualMaintenancePlanData.json). Manuales publicados en la biblioteca de la app.", "—"),
    ("Tableros eléctricos, tanques de agua, chumaceras y motobombas.", "RESUELTO", "Agregados al plan (LF-151 a LF-155 tanques, LF-161 a LF-164 motobombas, LF-191 a LF-194 tableros, LF-201 a LF-203 chumaceras) con su manual MAN-LF-06.", "—"),
    ("Firmas físicas de aprobación (portada y manuales).", "PENDIENTE", "El contenido está aprobado. Falta firmar la portada del PRGMAT01 v06 y la tabla de aprobación de cada manual.", "Mantenimiento / Coordinación / Gerencia"),
]


def hoja_pendientes(wb):
    ws = wb.create_sheet("12 PENDIENTES")
    cols = ["#", "Pendiente", "Estado", "Cómo quedó resuelto", "Responsable y plazo"]
    titulo(ws, "PENDIENTES — ESTADO AL 2026-10-06", "Lo que requiere presencia en campo quedó convertido en la tarea LF-100, con fecha en el cronograma y formato prellenado.", len(cols))
    hdr(ws, 5, cols)
    color = {"RESUELTO": "C6EFCE", "PROGRAMADO": "FFF4C2", "PENDIENTE": "F8CBAD"}
    for i, (a, est, b, c) in enumerate(PENDIENTES, 1):
        r = 5 + i
        celda(ws, r, 1, i, center=True); celda(ws, r, 2, a); celda(ws, r, 3, est, bold=True, center=True, fill=color[est]); celda(ws, r, 4, b); celda(ws, r, 5, c)
    for i, w in enumerate([4, 44, 13, 80, 28], 1): ws.column_dimensions[get_column_letter(i)].width = w

# ───────────────────────── Excel del plan ─────────────────────────
thin = Side(style="thin", color="9AA8B8"); BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(ws, row, values, fill=NAVY):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v); c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=fill); c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = BORDE
def titulo(ws, txt, sub, ncols):
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=txt); c.font = Font(bold=True, color="FFFFFF", size=13); c.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols); ws.cell(row=3, column=1).fill = PatternFill("solid", fgColor=ORANGE)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ncols)
    c = ws.cell(row=4, column=1, value=sub); c.font = Font(size=9, color="44546A"); c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[4].height = 30
def celda(ws, r, c, v, wrap=True, bold=False, fill=None, size=9, center=False):
    x = ws.cell(row=r, column=c, value=v); x.font = Font(size=size, bold=bold); x.border = BORDE
    x.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if center else "left")
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x

def generar_excel():
    wb = Workbook()
    # 00 PORTADA
    ws = wb.active; ws.title = "00 PORTADA"
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 90
    titulo(ws, "PLAN AM — G-GRANJA LA FE · v06 APROBADO", "Reemplaza la hoja «06.2 PLAN LA FE» del PRGMAT01 v05 (incorporado al PRGMAT01 v06). Mismo esquema de columnas y mismos pasos de ejecución (PROMAT01 v02 sección 4) que el Plan AM de Planta Incubant.", 2)
    filas = [
        ("Empresa", EMPRESA), ("NIT", NIT), ("Sede", "G-GRANJA LA FE · Hispania, Antioquia"),
        ("Documento base", "PRGMAT01 Programa de mantenimiento preventivo v05 — hoja 06.2"), ("Versión", "06 · APROBADO (contenido aprobado por Mantenimiento el 2026-10-06; firmas físicas abajo)"), ("Fecha", HOY),
        ("Composición de la granja", "5 módulos · 24 galpones: 3 módulos de producción con 4 galpones cada uno (M2, M3, M4) y 2 módulos de levante con 6 galpones cada uno (M1, M5)."),
        ("Equipos de producción (Big Herdsman)", "Banda recolectora de huevo · sistema de apertura y cierre de nidos · alimentación (silos, tolvas, canal y cadena) · antipercheo con sistema de agua (niples) · iluminación inteligente."),
        ("Equipos de servicio", "Tableros eléctricos (24 de galpón + general, transferencia y bombeo) · motobombas · 11 tanques de agua · chumaceras de banda, nidos, alimentación y transportadores · planta eléctrica y ATS."),
        ("Qué cambia frente a la v05", "Se agregan los sistemas que faltaban (antipercheo e iluminación inteligente); la alimentación de producción pasa de «sinfines» a canal y cadena con tolvas (equipo real); se separan las tareas de levante; se agregan limpieza de silos, verificación de reguladores y lavado/overhaul por vacío sanitario; se agregan planes propios para tableros eléctricos, motobombas, chumaceras e inspección estructural de tanques; tarea única de levantamiento de campo LF-100; cada tarea trae pasos técnicos y manual de referencia. Códigos LF-100 a LF-221 (los LF-001 a LF-030 de la v05 quedan retirados; columna «Reemplaza a» para trazabilidad)."),
        ("Normas", "Res. ICA 3651 de 2014 (An. núm. 1.9 mantenimiento, núm. 1.3 limpieza y desinfección, Art. 4.2.12-4.2.13 POE y registros ≥ 1 año) · RETIE · Res. 4272 de 2021 (trabajo en alturas) · Res. 0491 de 2020 (espacios confinados)."),
        ("Cómo se lee el cronograma", "Ciclo de 4 semanas, un módulo por semana: S1 → módulo 2, S2 → módulo 3, S3 → módulo 4, S4 → levante (módulos 1 y 5). Todas las tareas con parada de un módulo se hacen la misma semana; el técnico no cruza entre módulos de producción en la misma semana (bioseguridad). En la celda: 2/3/4 = módulo de producción, L = levante, ● = todos. Diarias = ronda; por vacío sanitario = se fija con la fecha real de vaciado de cada galpón. Las tareas semanales sin parada las hace el auxiliar asignado a cada módulo o, si es uno solo, en el orden del día levante → producción, con ducha y cambio de ropa entre módulos."),
        ("Trimestrales", "Caen en la primera semana del trimestre en que se visita el módulo; por eso el intervalo real entre una y otra va de 12 a 16 semanas (tolerancia ± 2 semanas)."),
    ]
    for i, (k, v) in enumerate(filas, 6):
        celda(ws, i, 1, k, bold=True, fill=GRIS); celda(ws, i, 2, v)
    r = 6 + len(filas) + 1
    celda(ws, r, 1, "APROBACIONES", bold=True, fill=GRIS); celda(ws, r, 2, "Este plan entra en vigor con las tres firmas.", bold=True)
    for j, k in enumerate(["Elaboró — Área de Mantenimiento", "Revisó — Coordinación de Granja", "Aprobó — Gerencia"], r + 1):
        celda(ws, j, 1, k, bold=True); celda(ws, j, 2, "Nombre: ______________________    Cargo: ______________________    Firma: ____________    Fecha: ________")

    # 05 INVENTARIO LA FE
    ws = wb.create_sheet("05 INVENTARIO LA FE")
    cols = ["Módulo", "Uso", "Galpón (código plan)", "Banda de huevo", "Nidos automáticos", "Alimentación", "Agua (niples)", "Antipercheo", "Iluminación inteligente", "Marca", "Observación"]
    titulo(ws, "INVENTARIO DE SISTEMAS POR GALPÓN — G-GRANJA LA FE", "Celdas «diligenciar»: modelo, n.º de serie y datos de placa de cada motor se toman en campo. Inventario de galpones del PRGMAT01 hoja 05: cada código de galpón doble (p. ej. G-10101 «Galpón 101-102 Piso 1») corresponde a un galpón del plan.", len(cols))
    hdr(ws, 5, cols)
    r = 6
    for mod, uso, gal, desc in MODULOS:
        for g in gal:
            prod = uso == "Producción"
            vals = [mod, uso, g, "Sí" if prod else "No", "Sí" if prod else "No",
                    "Silo + tolva + canal y cadena" if prod else "Por confirmar", "Sí", "Sí" if prod else "No", "Sí" if prod else "Por confirmar",
                    "BIG HERDSMAN" if prod else "Por confirmar", desc]
            for i, v in enumerate(vals, 1): celda(ws, r, i, v, fill="FFF4C2" if v == "Por confirmar" else None)
            r += 1
    for i, w in enumerate([10, 11, 12, 10, 10, 18, 9, 11, 13, 15, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"

    # 06.2 PLAN LA FE
    ws = wb.create_sheet("06.2 PLAN LA FE")
    cols = ["Código", "Sistema", "Clase de equipo", "Equipos que aplican", "Cant.", "Descripción del procedimiento", "Tipo", "Periodicidad", "¿Parada?", "Dur. (h/eq)", "Responsable (rol)", "Criterio de aceptación", "Sustento de la periodicidad", "Registro / evidencia", "Manual", "Reemplaza a (v05)"]
    titulo(ws, "PLAN DE MANTENIMIENTO — G-GRANJA LA FE (v06)", "Res. ICA 3651 de 2014 — granja avícola de postura y/o levante biosegura  ·  POE núm. 1.9 (mantenimiento) · POE núm. 1.3 (limpieza y desinfección) · Art. 4.2.12 y 4.2.13 (POE y registros ≥ 1 año)", len(cols))
    ws.cell(row=3, column=1, value="LEYENDA:   ■ Rojo = tarea crítica de seguridad (su omisión causa pérdida de aves, del lote o un accidente) · Pasos de cada actividad: hoja 06.2-IND y documento «Indicaciones de cada actividad»").font = Font(bold=True, color="FFFFFF", size=9)
    hdr(ws, 5, cols)
    for i, t in enumerate(TAREAS, 6):
        n = len([e for e in t["equipos"] if re.match(r"^[GS]\d", e)]) or "—"
        vals = [t["code"], t["system"], t["eq_class"], ", ".join(t["equipos"]), n, t["desc"], t["tipo"], t["freq"], t["parada"], t["dur"], t["resp"], t["crit"], t["sust"], t["evid"], t["manual"], t["reemplaza"]]
        for j, v in enumerate(vals, 1):
            c = celda(ws, i, j, v, center=j in (5, 9, 10))
            if t["cs"] and j == 1: c.fill = PatternFill("solid", fgColor="C00000"); c.font = Font(bold=True, color="FFFFFF", size=9)
    for i, w in enumerate([11, 13, 24, 30, 6, 50, 13, 14, 8, 8, 16, 38, 30, 16, 18, 11], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B6"; ws.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{5 + len(TAREAS)}"

    # 06.2-IND INDICACIONES
    ws = wb.create_sheet("06.2-IND INDICACIONES")
    cols = ["Código", "Actividad", "Paso PROMAT01", "Qué hacer", "Responsable", "Registro"]
    titulo(ws, "INDICACIONES DE CADA ACTIVIDAD — PASOS PROMAT01 v02 SECCIÓN 4 + PASOS TÉCNICOS DEL MANUAL", "Igual al recuadro que IncubApp abre al hacer clic en una actividad del Plan AM de planta; el paso 2 se detalla con los pasos técnicos del manual del equipo.", len(cols))
    hdr(ws, 5, cols); r = 6
    for t in TAREAS:
        first = r
        for titulo_p, det, resp, reg in pasos_promat(t):
            texto = " ".join(det)
            if titulo_p.startswith("2."):
                texto += "\n\nPASOS TÉCNICOS:\n" + "\n".join(f"{k}. {s}" for k, s in enumerate(t["pasos"], 1))
                if t["seg"]: texto += "\n\nSEGURIDAD:\n" + "\n".join("• " + s for s in t["seg"])
                if t["herr"]: texto += "\n\nHERRAMIENTAS: " + ", ".join(t["herr"])
            celda(ws, r, 1, t["code"], bold=True); celda(ws, r, 2, t["desc"]); celda(ws, r, 3, titulo_p, bold=True, fill=GRIS)
            celda(ws, r, 4, texto); celda(ws, r, 5, resp); celda(ws, r, 6, reg)
            r += 1
        ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)
        ws.merge_cells(start_row=first, start_column=2, end_row=r - 1, end_column=2)
    for i, w in enumerate([11, 34, 22, 90, 18, 26], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C6"

    # 07 CRONOGRAMA
    ws = wb.create_sheet("07 CRONOGRAMA 52 SEM")
    fijos = ["Código", "Sistema", "Clase de equipo", "Actividad (resumen)", "Periodicidad", "Lote / rotación"]
    ncols = len(fijos) + 52
    titulo(ws, "CRONOGRAMA ANUAL — 52 SEMANAS — G-GRANJA LA FE", "Ciclo de 4 semanas por módulo: S1 = módulo 2 · S2 = módulo 3 · S3 = módulo 4 · S4 = levante (M1 y M5). En la celda: 2/3/4 = módulo de producción, L = levante, ● = todos los equipos. Diarias: ronda. Por vacío sanitario / entre lotes: se programa con la fecha real de vaciado de cada galpón.", ncols)
    meses = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
    celda(ws, 5, 1, "MES →", bold=True)
    for w in range(1, 53):
        m = min(11, int((w - 1) * 12 / 52))
        c = ws.cell(row=5, column=len(fijos) + w, value=meses[m] if (w == 1 or min(11, int((w - 2) * 12 / 52)) != m) else None)
        c.font = Font(bold=True, size=8)
    hdr(ws, 6, fijos + [str(w) for w in range(1, 53)])
    ws.cell(row=7, column=1, value="Módulo de la semana").font = Font(bold=True, size=8)
    for w in range(1, 53):
        m = CICLO[(w - 1) % 4]
        celda(ws, 7, len(fijos) + w, {"M2": "M2", "M3": "M3", "M4": "M4", "LEV": "LEV"}[m], size=7, center=True, fill="DDE6F0")
    for i, t in enumerate(TAREAS, 8):
        cr = cronograma(t)
        rot = t["rot"]
        for j, v in enumerate([t["code"], t["system"], t["eq_class"], t["desc"][:90] + ("…" if len(t["desc"]) > 90 else ""), t["freq"], rot], 1):
            celda(ws, i, j, v, wrap=False, size=8)
        for w in range(1, 53):
            v = cr.get(w)
            c = celda(ws, i, len(fijos) + w, v, size=8, center=True, fill=("E0740A" if v else None))
            if v: c.font = Font(size=8, bold=True, color="FFFFFF")
        if not cr:
            c = ws.cell(row=i, column=len(fijos) + 1, value="Ronda diaria" if t["freq"] == "Diaria" else "Se programa con la fecha real de vaciado del galpón")
            c.font = Font(italic=True, size=8, color="44546A")
    for i, w in enumerate([10, 13, 26, 46, 16, 22], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for w in range(1, 53): ws.column_dimensions[get_column_letter(len(fijos) + w)].width = 3.6
    ws.freeze_panes = ws.cell(row=8, column=len(fijos) + 1)

    # 08 CARGA
    ws = wb.create_sheet("08 CARGA")
    cols = ["Periodicidad", "N.º de tareas", "Equipos-intervención / año", "Horas-hombre / año (est.)", "Nota"]
    titulo(ws, "DIMENSIONAMIENTO DE CARGA DE TRABAJO ANUAL — G-GRANJA LA FE", "Estimación de ingeniería para dimensionar la plantilla antes de aprobar. Vacío sanitario ≈ 1 por galpón al año; entre lotes ≈ 1 por galpón al año.", len(cols))
    hdr(ws, 5, cols)
    VECES = {"Diaria": 365, "Semanal": 52, "Mensual": 12, "Trimestral": 4, "Semestral": 2, "Anual": 1}
    agg = {}
    for t in TAREAS:
        f = t["freq"]; veces = VECES.get(f, 1); k = f
        n = len(t["equipos"])
        a = agg.setdefault(k, [0, 0, 0.0]); a[0] += 1; a[1] += n * veces; a[2] += n * veces * t["dur"]
    r = 6; tot = 0
    for k, (nt, ni, hh) in sorted(agg.items()):
        celda(ws, r, 1, k); celda(ws, r, 2, nt, center=True); celda(ws, r, 3, ni, center=True); celda(ws, r, 4, round(hh, 1), center=True)
        celda(ws, r, 5, "Rondas del operario galponero: incluidas en su jornada" if k == "Diaria" else ""); tot += hh if k != "Diaria" else 0; r += 1
    celda(ws, r, 1, "TOTAL mantenimiento (sin rondas diarias)", bold=True); celda(ws, r, 4, round(tot, 1), bold=True, center=True)
    celda(ws, r, 5, f"≈ {round(tot / 2000, 2)} personas de tiempo completo (2.000 h/año), sin correctivos ni desplazamientos.", bold=True)
    for i, w in enumerate([34, 12, 22, 22, 70], 1): ws.column_dimensions[get_column_letter(i)].width = w

    # 09 REGISTRO EJECUCION
    ws = wb.create_sheet("09 REGISTRO EJECUCION")
    cols = ["Fecha", "Código de tarea", "Sede", "Equipo (código)", "Descripción de lo ejecutado", "Ejecutor (nombre)", "Hora inicio", "Hora fin", "Resultado", "Hallazgo y acción correctiva", "Firma / N° OT IncubApp"]
    titulo(ws, "FORMATO DE REGISTRO DE EJECUCIÓN", "Evidencia de cada intervención. Conservar mínimo 1 año en la instalación (Res. ICA 3651/2014 Art. 4.2.13).", len(cols))
    hdr(ws, 5, cols)
    for r in range(6, 66):
        for c in range(1, len(cols) + 1): celda(ws, r, c, "G-GRANJA LA FE" if c == 3 else None)
    for i, w in enumerate([11, 12, 15, 14, 40, 20, 9, 9, 14, 34, 18], 1): ws.column_dimensions[get_column_letter(i)].width = w

    hoja_levantamiento(wb)
    hoja_pendientes(wb)

    ruta = os.path.join(BASE, "PRGMAT01 06.2 Plan AM Granja La Fe (v06 aprobado).xlsx")
    wb.save(ruta); return ruta

# ───────────────────────── JSON (formato annualMaintenancePlanData.tasks) ─────────────────────────
def generar_json():
    out = []
    for t in TAREAS:
        eq = [e for e in t["equipos"] if re.match(r"^[GS]\d", e)]
        out.append({
            "code": t["code"], "sede": "GRANJA LA FE", "system": t["system"].upper(), "equipmentClass": t["eq_class"],
            "applyingEquipment": ", ".join(t["equipos"]), "count": len(eq) or 1, "description": t["desc"], "type": t["tipo"],
            "frequency": t["freq"], "shutdown": t["parada"], "duration": str(t["dur"]), "responsible": t["resp"],
            "acceptanceCriteria": t["crit"], "frequencyRationale": t["sust"], "evidenceFormat": t["evid"],
            "isCriticalSecurity": t["cs"], "isBiosecurity": t["bio"],
            "cronograma": {"rotacion": t["rot"], "weeks": semanas(t)},
            "manual": t["manual"], "replaces": t["reemplaza"], "technicalSteps": t["pasos"], "safety": t["seg"], "tools": t["herr"],
        })
    os.makedirs(os.path.join(BASE, "datos"), exist_ok=True)
    ruta = os.path.join(BASE, "datos", "plan_am_la_fe.json")
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump({"metadata": {"generatedAt": HOY, "source": "PRGMAT01 v06 hoja 06.2 — aprobado 2026-10-06", "totalTasks": len(out)}, "tasks": out}, f, ensure_ascii=False, indent=1)
    return ruta

# ───────────────────────── Carta al fabricante (pendiente 1) ─────────────────────────
def generar_carta():
    doc = Document(); set_base(doc)
    doc.add_paragraph(f"Hispania (Antioquia, Colombia), {HOY}")
    doc.add_paragraph("Para: Qingdao Big Herdsman Machinery Co., Ltd. — bigherdsman@bigherdsman.com\nDe: Área de Mantenimiento, " + EMPRESA + f" (NIT {NIT})\nAsunto: Solicitud de manuales de operación y mantenimiento / Request for operation and maintenance manuals")
    h(doc, "Español", 2)
    doc.add_paragraph("Respetados señores: en nuestra granja de reproductoras G-GRANJA LA FE operamos equipos Big Herdsman en 12 galpones de producción: banda recolectora de huevo, nidos automáticos con apertura y cierre, sistema de alimentación (silos, transportador, tolvas, canal y cadena), sistema antipercheo con bebederos de niple e iluminación inteligente. Les solicitamos, en español o en inglés, los manuales de instalación, operación y mantenimiento de esos equipos, con el plan de lubricación, los límites de desgaste (elongación de cadena, banda), las tablas de caudal de niples, los datos del energizador antipercheo y del controlador de iluminación, y la lista de repuestos recomendados. Los usaremos para validar nuestro plan de mantenimiento (Resolución ICA 3651 de 2014). Gracias.")
    h(doc, "English", 2)
    doc.add_paragraph("Dear Sirs: at our broiler breeder farm G-GRANJA LA FE we operate Big Herdsman equipment in 12 production houses: egg collection belt, automatic nests with opening/closing mechanism, feeding system (silos, feed conveyor, hoppers, chain feeder), anti-perching system on nipple drinker lines, and smart lighting. We kindly request, in Spanish or English, the installation, operation and maintenance manuals for this equipment, including the lubrication schedule, wear limits (chain elongation, belt), nipple flow tables, anti-perching energizer and lighting controller data, and the recommended spare parts list. They will be used to validate our maintenance plan (ICA Resolution 3651/2014). Thank you.")
    doc.add_paragraph("\n\n______________________________\nÁrea de Mantenimiento — " + EMPRESA)
    ruta = os.path.join(BASE, "Carta solicitud manuales Big Herdsman.docx")
    doc.save(ruta); return ruta


if __name__ == "__main__":
    os.makedirs(os.path.join(BASE, "Manuales"), exist_ok=True)
    for m in MANUALES: print(generar_manual(m))
    print(generar_indicaciones()); print(generar_excel()); print(generar_json()); print(generar_carta())
